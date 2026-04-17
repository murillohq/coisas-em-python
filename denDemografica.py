import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, Reference

def populacao_relativa(populacao, area):
    if area <= 0:
        raise ValueError("A área deve ser maior que zero.")
    if populacao < 0:
        raise ValueError("A população não pode ser negativa.")
    return populacao / area

# Dados
dados = [
    ("Brasil",  211700000, 8500000),
    ("Goiás",   7056495,   340203),
    ("Jataí",   105729,    7174),
]

# Print no terminal
for nome, pop, area in dados:
    resultado = populacao_relativa(pop, area)
    print(f"{nome}: {resultado:.2f} hab/km²")

# Excel
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Densidade"

# Cabeçalho
cabecalho = ["Local", "População", "Área (km²)", "Densidade (hab/km²)"]
ws.append(cabecalho)

estilo_cab = Font(bold=True, color="FFFFFF")
fundo_cab  = PatternFill("solid", fgColor="2E86AB")

for col, _ in enumerate(cabecalho, start=1):
    cel = ws.cell(row=1, column=col)
    cel.font = estilo_cab
    cel.fill = fundo_cab
    cel.alignment = Alignment(horizontal="center")

# Dados nas linhas
for nome, pop, area in dados:
    densidade = populacao_relativa(pop, area)
    ws.append([nome, pop, area, round(densidade, 2)])

# Largura das colunas
ws.column_dimensions["A"].width = 12
ws.column_dimensions["B"].width = 18
ws.column_dimensions["C"].width = 14
ws.column_dimensions["D"].width = 22

# Gráfico de barras
chart = BarChart()
chart.type = "col"
chart.title = "Densidade Demográfica (hab/km²)"
chart.y_axis.title = "hab/km²"
chart.x_axis.title = "Local"
chart.style = 10
chart.width = 15
chart.height = 10

# Referência dos dados (coluna D = densidade)
dados_ref = Reference(ws, min_col=4, min_row=1, max_row=4)
locais_ref = Reference(ws, min_col=1, min_row=2, max_row=4)

chart.add_data(dados_ref, titles_from_data=True)
chart.set_categories(locais_ref)

ws.add_chart(chart, "F2")

# Salvar
wb.save("densidade_demografica.xlsx")
print("\nArquivo salvo: densidade_demografica.xlsx")